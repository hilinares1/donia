import string

from odoo import api, fields, models, exceptions
import logging
import base64
import openpyxl
import tempfile
import string
import random


_logger = logging.getLogger()


class LoadExcelFile(models.TransientModel):
    _name = 'wizard.load_file'

    serial_file = fields.Binary(string="Cargar Archivo")

    def generate_random_name(self):
        letters = string.ascii_lowercase
        return ''.join(random.choice(letters) for i in range(10))

    def load_excel_file(self):
        # raise exceptions.Warning("Probando el wizard")

        path = ''
        name_of_file = self.generate_random_name()

        #temp = tempfile.NamedTemporaryFile(prefix="modified_file_" + name_of_file + '')
        with tempfile.NamedTemporaryFile(prefix="modified_file_" + name_of_file, suffix='.xlsx', delete=False) as tmp:
            tmp.write(base64.decodebytes(self.serial_file))

        # file_decoded = base64.decodebytes(self.serial_file)
        # file_manifest = open("/tmp/prueba.xlsx", "wb")
        # file_manifest.write(file_decoded)
        # file_manifest.close()

        xfile = openpyxl.load_workbook(tmp.name)
        sheet = xfile.worksheets[0]
        # sheet.cell(1, 6).value = 'hello world'

        row_count = sheet.max_row
        column_count = sheet.max_column

        row_with_value = 0

        for row in range(1, row_count + 1):
            if not sheet.cell(row, 1).value == None:
                row_with_value = row
            else:

                if sheet.cell(row, 4).value == None or sheet.cell(row, 4).value == sheet.cell(row_with_value, 4).value:

                    sheet.cell(row_with_value, 5).value = sheet.cell(row_with_value, 5).value + ',' + sheet.cell(row,
                                                                                                                 5).value
                    sheet.cell(row, 5).value = ""
                    sheet.cell(row, 1).value = 'delete'
                else:
                    row_with_value = row

        # deleting rows
        #         for row1 in range(1, row_count):
        #             if sheet.cell(row1, 1).value == 'delete':
        #                 sheet.delete_rows(row1, 1)

        #             for col in range(1, column_count + 1):
        #                 if sheet.cell(row, 1)

        xfile.save(tmp.name)
        self.clean_file_deleted_row(tmp)

        with open(tmp.name, 'rb') as r:
            file = base64.b64encode(r.read())

        att_vals = {
            'name': "modified_file_" + name_of_file + '.xlsx',
            'type': 'binary',
            'datas': file,
        }

        attachment_id = self.env['ir.attachment'].create(att_vals)
        self.env.cr.commit()

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/{}?download=true'.format(attachment_id.id, ),
            'target': 'self',
        }

    def clean_file_deleted_row(self, tmp):
        xfile = openpyxl.load_workbook(tmp.name)
        sheet = xfile.worksheets[0]
        # sheet.cell(1, 6).value = 'hello world'

        row_count = sheet.max_row
        column_count = sheet.max_column

        for row in range(1, row_count + 1):
            for row1 in range(row, row_count + 1):
                if sheet.cell(row1, 1).value == 'delete':
                    sheet.delete_rows(row1, 1)

        xfile.save(tmp.name)

